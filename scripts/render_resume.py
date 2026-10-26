#!/usr/bin/env python3
"""Render a resume JSON file to a one-page PDF and a Word file in the house layout.

Usage:
  python3 render_resume.py resume.json --out "First Last Resume"
  -> writes "First Last Resume.pdf" and "First Last Resume.docx"
  python3 render_resume.py resume.json --out "First Last Resume (preview)" --preview
  -> writes a PDF only; placeholders are allowed and highlighted

Placeholders: write facts the user still has to give as [[...]], e.g.
  "raising activation from [[X%]] to [[Y%]]" or "dates": "[[Month Year]] – [[Month Year]]".
They render as [X%] with a yellow highlight. A final render (no --preview) refuses to run
while any placeholder is left, and lists them.

Layout: centered name, contact line with links, section headings with a full-width rule,
organization + location on one line, title + dates on the next, round bullets.
Serif font (Times in the PDF, Times New Roman in Word). The font size steps down from
11 pt to 10 pt to fit one page; if it still doesn't fit, the script says so.

JSON shape (see templates/resume.example.json):
{
  "name": "First Last",
  "pronunciation": "(optional phonetic line)",
  "contact": ["555-123-4567", "me@example.com", "linkedin.com/in/me", "github.com/me"],
  "sections": [
    {"title": "Summary", "text": "Two or three lines."},
    {"title": "Education", "entries": [{"org": "", "location": "", "title": "", "dates": "",
                                         "lines": ["Relevant Coursework: ..."], "bullets": []}]},
    {"title": "Skills", "lines": ["Languages: Python, SQL", "Tools: Figma, Jira"]}
  ]
}
In "lines", text before the first colon is set in bold (e.g. "Relevant Coursework:").
"""
import argparse
import json
import re
import sys

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

PAGE_W, PAGE_H = letter
MARGIN = 0.55 * inch
TEXT_W = PAGE_W - 2 * MARGIN
FONTS = {"r": "Times-Roman", "b": "Times-Bold", "i": "Times-Italic"}
BULLET_INDENT = 14
SIZES = [11, 10.5, 10]
PH = re.compile(r"\[\[(.+?)\]\]")
NBSP = "\u00a0"


def shown(s):
    """Text as displayed: [[X]] -> [X]."""
    return PH.sub(lambda m: "[" + m.group(1) + "]", str(s)).replace(NBSP, " ")


def keep_ph_together(s):
    """Stop placeholders from wrapping across lines."""
    return PH.sub(lambda m: "[[" + m.group(1).replace(" ", NBSP) + "]]", str(s))


def placeholders(obj):
    if isinstance(obj, str):
        return PH.findall(obj)
    if isinstance(obj, list):
        return [p for x in obj for p in placeholders(x)]
    if isinstance(obj, dict):
        return [p for v in obj.values() for p in placeholders(v)]
    return []


def link_for(item):
    s = item.strip()
    if PH.search(s):
        return None
    if "@" in s and " " not in s:
        return "mailto:" + s
    if re.match(r"^(www\.)?[\w-]+(\.[\w-]+)+(/\S*)?$", s) and not re.match(r"^[\d\-\s().+]+$", s):
        return "https://" + s
    return None


def width_of(s, font, size):
    return stringWidth(shown(s), font, size)


def wrap(text, font, size, width):
    words, lines, cur = [w for w in re.split(r" +", keep_ph_together(text).strip()) if w], [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if width_of(trial, font, size) <= width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


class PdfLayout:
    """Lays out the resume. With draw=False it only measures height."""

    def __init__(self, data, size, c=None):
        self.d, self.s, self.c = data, size, c
        self.lead = size * 1.18
        self.y = PAGE_H - MARGIN

    def text(self, x, s, font, size=None, right=False, center=False):
        size = size or self.s
        if not self.c:
            return
        w = width_of(s, FONTS[font], size)
        x = x - w if right else x - w / 2 if center else x
        self.c.setFont(FONTS[font], size)
        for seg in re.split(r"(\[\[.+?\]\])", str(s)):
            if not seg:
                continue
            is_ph = PH.fullmatch(seg) is not None
            t = shown(seg)
            sw = stringWidth(t, FONTS[font], size)
            if is_ph:
                self.c.setFillColorRGB(1, 0.92, 0.45)
                self.c.rect(x - 1, self.y - size * 0.22, sw + 2, size * 1.02, stroke=0, fill=1)
                self.c.setFillColorRGB(0, 0, 0)
            self.c.drawString(x, self.y, t)
            x += sw

    def down(self, h):
        self.y -= h

    def labeled_line(self, line, x, width):
        m = re.match(r"^([^:]{1,40}:)\s*(.*)$", line)
        if not m:
            for ln in wrap(line, FONTS["r"], self.s, width):
                self.down(self.lead)
                self.text(x, ln, "r")
            return
        label, rest = m.group(1), m.group(2)
        lw = width_of(label + " ", FONTS["b"], self.s)
        first = wrap(rest, FONTS["r"], self.s, width - lw)
        self.down(self.lead)
        self.text(x, label, "b")
        if first:
            self.text(x + lw, first[0], "r")
            rest_words = [w for w in re.split(r" +", keep_ph_together(rest).strip()) if w]
            remainder = " ".join(rest_words[len(first[0].split(" ")):])
            for ln in wrap(remainder, FONTS["r"], self.s, width):
                self.down(self.lead)
                self.text(x, ln, "r")

    def run(self):
        d, s = self.d, self.s
        # name
        self.down(s * 2.0)
        self.text(PAGE_W / 2, d.get("name", ""), "b", size=s * 2.0, center=True)
        if d.get("pronunciation"):
            self.down(self.lead)
            self.text(PAGE_W / 2, d["pronunciation"], "i", center=True)
        # contact line with links
        items = [str(i) for i in d.get("contact", []) if str(i).strip()]
        if items:
            self.down(self.lead * 1.25)
            sep = "  |  "
            full = sep.join(items)
            x = PAGE_W / 2 - width_of(full, FONTS["r"], s) / 2
            for k, it in enumerate(items):
                w = width_of(it, FONTS["r"], s)
                self.text(x, it, "r")
                url = link_for(it)
                if self.c and url:
                    self.c.linkURL(url, (x, self.y - 2, x + w, self.y + s), relative=0)
                x += w
                if k < len(items) - 1:
                    self.text(x, sep, "r")
                    x += stringWidth(sep, FONTS["r"], s)
        for sec in d.get("sections", []):
            self.section(sec)
        return PAGE_H - MARGIN - self.y

    def section(self, sec):
        s = self.s
        self.down(s * 1.15)
        self.down(s * 1.2)
        self.text(MARGIN, sec.get("title", "").upper(), "b", size=s * 1.1)
        if self.c:
            self.c.setLineWidth(0.6)
            self.c.line(MARGIN, self.y - 3, PAGE_W - MARGIN, self.y - 3)
        self.down(3)
        if sec.get("text"):
            for ln in wrap(sec["text"], FONTS["r"], s, TEXT_W):
                self.down(self.lead)
                self.text(MARGIN, ln, "r")
        for line in sec.get("lines", []):
            self.labeled_line(line, MARGIN, TEXT_W)
        for i, e in enumerate(sec.get("entries", [])):
            if i:
                self.down(s * 0.35)
            self.entry(e)

    def entry(self, e):
        s = self.s
        self.down(self.lead)
        self.text(MARGIN, e.get("org", ""), "b")
        self.text(PAGE_W - MARGIN, e.get("location", ""), "r", right=True)
        if e.get("title") or e.get("dates"):
            self.down(self.lead)
            self.text(MARGIN, e.get("title", ""), "i")
            self.text(PAGE_W - MARGIN, e.get("dates", ""), "i", right=True)
        for line in e.get("lines", []):
            self.labeled_line(line, MARGIN + BULLET_INDENT, TEXT_W - BULLET_INDENT)
        for b in e.get("bullets", []):
            lines = wrap(b, FONTS["r"], s, TEXT_W - BULLET_INDENT - 6)
            for k, ln in enumerate(lines):
                self.down(self.lead)
                if k == 0:
                    self.text(MARGIN + 4, "•", "r")
                self.text(MARGIN + BULLET_INDENT, ln, "r")


def pick_size(data):
    avail = PAGE_H - 2 * MARGIN
    for size in SIZES:
        if PdfLayout(data, size).run() <= avail:
            return size, True
    return SIZES[-1], False


def build_pdf(data, size, path):
    c = canvas.Canvas(path, pagesize=letter)
    c.setTitle(f"{data.get('name', '')} Resume")
    c.setAuthor(data.get("name", ""))
    PdfLayout(data, size, c).run()
    c.showPage()
    c.save()


def build_docx(data, size, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.55))
    right_tab = Inches(8.5 - 1.1)

    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    def para(align=None, before=0, indent=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(before)
        if indent is not None:
            p.paragraph_format.left_indent = indent
        return p

    from docx.enum.text import WD_COLOR_INDEX

    def run(p, text, bold=False, italic=False, sz=None):
        r = None
        for seg in re.split(r"(\[\[.+?\]\])", str(text)):
            if not seg:
                continue
            r = p.add_run(shown(seg))
            r.bold, r.italic = bold, italic
            if sz:
                r.font.size = Pt(sz)
            if PH.fullmatch(seg):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return r

    def hyperlink(p, url, text):
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        p._p.append(h)

    def two_sided(left, right, style):
        p = para()
        p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
        run(p, left, bold=style == "b", italic=style == "i")
        if right:
            run(p, "\t" + right, italic=style == "i")
        return p

    def labeled(line, indent=None):
        p = para(indent=indent)
        m = re.match(r"^([^:]{1,40}:)\s*(.*)$", line)
        if m:
            run(p, m.group(1) + " ", bold=True)
            run(p, m.group(2))
        else:
            run(p, line)

    p = para(WD_ALIGN_PARAGRAPH.CENTER)
    run(p, data.get("name", ""), bold=True, sz=size * 2)
    if data.get("pronunciation"):
        run(para(WD_ALIGN_PARAGRAPH.CENTER), data["pronunciation"], italic=True)
    items = [str(i) for i in data.get("contact", []) if str(i).strip()]
    if items:
        p = para(WD_ALIGN_PARAGRAPH.CENTER, before=2)
        for k, it in enumerate(items):
            url = link_for(it)
            hyperlink(p, url, it) if url else run(p, it)
            if k < len(items) - 1:
                run(p, "  |  ")

    for s in data.get("sections", []):
        h = para(before=size * 0.9)
        run(h, s.get("title", "").upper(), bold=True, sz=size * 1.1)
        pPr = h._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)
        if s.get("text"):
            run(para(), s["text"])
        for line in s.get("lines", []):
            labeled(line)
        for i, e in enumerate(s.get("entries", [])):
            p = two_sided(e.get("org", ""), e.get("location", ""), "b")
            if i:
                p.paragraph_format.space_before = Pt(size * 0.35)
            if e.get("title") or e.get("dates"):
                two_sided(e.get("title", ""), e.get("dates", ""), "i")
            for line in e.get("lines", []):
                labeled(line, indent=Inches(0.2))
            for b in e.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(0)
                run(bp, b)
    doc.core_properties.title = f"{data.get('name', '')} Resume"
    doc.core_properties.author = data.get("name", "")
    doc.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("resume_json")
    ap.add_argument("--out", required=True, help="output path without extension")
    ap.add_argument("--pdf-only", action="store_true")
    ap.add_argument("--preview", action="store_true", help="draft: allow [[placeholders]], PDF only")
    a = ap.parse_args()
    with open(a.resume_json, encoding="utf-8") as f:
        data = json.load(f)
    left = placeholders(data)
    if left and not a.preview:
        print("Placeholders still to fill (use --preview to render a draft):", file=sys.stderr)
        for p in left:
            print("  [" + p + "]", file=sys.stderr)
        sys.exit(3)
    size, fits = pick_size(data)
    build_pdf(data, size, a.out + ".pdf")
    print(f"PDF: {a.out}.pdf ({size} pt)")
    if left:
        print(f"{len(left)} placeholder(s) highlighted: " + ", ".join("[" + p + "]" for p in left))
    if not (a.pdf_only or a.preview):
        build_docx(data, size, a.out + ".docx")
        print(f"Word: {a.out}.docx")
    if not fits:
        print("WARNING: content is longer than one page even at 10 pt. Cut or shorten bullets.", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
